from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "00_SYSTEM_MODULES_AND_CHANGE_STANDARD.md"
SOURCE_TEXT = SOURCE.read_text(encoding="utf-8")


def source_header_value(label: str) -> str:
    match = re.search(rf"^> {re.escape(label)}：`([^`]+)`", SOURCE_TEXT, re.MULTILINE)
    if not match:
        raise RuntimeError(f"Missing {label} in {SOURCE}")
    return match.group(1)


DOCUMENT_VERSION = source_header_value("文档版本")
DATE = source_header_value("生效日期")
BASELINE_COMMIT = source_header_value("代码盘点基线")
OUTPUT = (
    ROOT
    / "output"
    / "docx"
    / f"LI3D_System_Maintenance_Manual_v{DOCUMENT_VERSION}.docx"
)
SCREENSHOT = ROOT / "docs" / "assets" / "LI3D_EDITOR_UI_BASELINE.png"

PRESET_NAME = "compact_reference_guide"
HEADER_PATTERN = "editorial_cover"
PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 80
CELL_MARGIN_START_END = 120

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
FONT_MONO = "Consolas"
INK = "182230"
NAVY = "183B56"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5C6673"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "B8C3CE"
WHITE = "FFFFFF"


def set_run_font(run, *, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 mono: bool = False) -> None:
    latin = FONT_MONO if mono else FONT_LATIN
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    rfonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(element, fill: str) -> None:
    if hasattr(element, "_tc"):
        props = element._tc.get_or_add_tcPr()
    elif hasattr(element, "_p"):
        props = element._p.get_or_add_pPr()
    elif element.tag.endswith("}tc"):
        props = element.get_or_add_tcPr()
    else:
        props = element.get_or_add_pPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border(paragraph, *, side: str = "left", color: str = BLUE,
                         size: int = 12, space: int = 8) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    pbdr.append(border)


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def prevent_table_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trpr.append(cant_split)


def set_cell_margins(cell, *, top: int = CELL_MARGIN_TOP_BOTTOM,
                     bottom: int = CELL_MARGIN_TOP_BOTTOM,
                     start: int = CELL_MARGIN_START_END,
                     end: int = CELL_MARGIN_START_END) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.find(qn("w:tcMar"))
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tcmar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tblind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths[index]))
            tcw.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_update_fields_on_open(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def create_numbering(doc: Document, *, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        numfmt = OxmlElement("w:numFmt")
        numfmt.set(qn("w:val"), "decimal" if ordered else "bullet")
        lvltext = OxmlElement("w:lvlText")
        lvltext.set(qn("w:val"), f"%{level + 1}." if ordered else "•")
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        position = 540 + level * 360
        text_position = 900 + level * 360
        tab.set(qn("w:pos"), str(text_position))
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(text_position))
        ind.set(qn("w:hanging"), str(text_position - position))
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.extend([tabs, ind, spacing])
        lvl.extend([start, numfmt, lvltext, suff, ppr])
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(max(0, min(level, 2))))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, numid])


def normalize_inline(text: str) -> str:
    return text.replace("  ", " ").strip()


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^\)]+\))")


def add_inline(paragraph, text: str, *, size: float = 11, color: str = INK,
               bold: bool = False, italic: bool = False) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color, bold=bold, italic=italic)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(8, size - 0.5), color=DARK_BLUE, bold=False, mono=True)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True, italic=italic)
        else:
            link_match = re.match(r"\[([^\]]+)\]\(([^\)]+)\)", token)
            label = link_match.group(1) if link_match else token
            run = paragraph.add_run(label)
            set_run_font(run, size=size, color=BLUE, bold=bold, italic=italic)
            run.font.underline = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.page_break_before = False


def configure_section(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("LI3D ENGINEERING  /  SYSTEM MAINTENANCE BASELINE")
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    set_paragraph_border(paragraph, side="bottom", color="D7DEE5", size=6, space=4)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(f"v{DOCUMENT_VERSION}   |   {DATE}   |   ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(paragraph, " PAGE ", "1")


def add_cover(doc: Document) -> None:
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LI3D ENGINEERING")
    set_run_font(run, size=10, color=BLUE, bold=True)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("系统模块与算法维护手册")
    set_run_font(run, size=29, color=NAVY, bold=True)
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("System Modules, Algorithms & Change-Control Manual")
    set_run_font(run, size=13.5, color=DARK_BLUE)
    p.paragraph_format.space_after = Pt(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("投影  /  UV 合成  /  局部重绘  /  图层  /  保存与版本历史")
    set_run_font(run, size=11, color=MUTED, bold=True)
    p.paragraph_format.space_after = Pt(70)

    metadata = [
        ("文档版本", DOCUMENT_VERSION),
        ("规范状态", "当前唯一维护准则 / Normative Baseline"),
        ("生效日期", DATE),
        ("代码盘点基线", BASELINE_COMMIT),
        ("适用对象", "两位维护者、Codex、后续开发与测试人员"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    for row, (label, value) in zip(table.rows, metadata):
        shade(row.cells[0], LIGHT_BLUE)
        add_inline(row.cells[0].paragraphs[0], label, size=9.5, color=NAVY, bold=True)
        add_inline(row.cells[1].paragraphs[0], value, size=9.5, color=INK)
    # Word accessibility checkers require a marked first row even for this
    # compact key/value document-control table.
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [1800, 7560])
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    p = doc.add_paragraph("文档控制", style="Heading 1")
    p.paragraph_format.page_break_before = False
    intro = doc.add_paragraph()
    shade(intro, CALLOUT)
    set_paragraph_border(intro, side="left", color=BLUE, size=16, space=10)
    add_inline(
        intro,
        "本文档既是系统模块目录，也是算法注册表和强制变更流程。任何行为修改必须先定位 M/UI/ALG ID，再建立 CHG 记录；历史方案不能覆盖本手册的当前真值。",
        size=10.5,
        color=INK,
        bold=True,
    )
    control_rows = [
        ("规范范围", "模块边界、界面调用、算法参数、图层语义、保存、版本、测试与变更控制"),
        ("真值原则", "代码是实现证据；本手册是维护规范。冲突必须建立变更记录后同时修正。"),
        ("算法版本", "规范版本与运行版本分离；没有持久化标记的算法明确登记为治理缺口。"),
        ("源码历史", "Git + CHANGELOG + docs/changes + ADR；工程 autosave 不能代替源码版本。"),
        ("默认修改上限", "普通 Bug 修复限制在一个大模块、通常不超过 5 个源文件或 300 行净变化。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "控制项"
    table.rows[0].cells[1].text = "当前规则"
    for cell in table.rows[0].cells:
        shade(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9, color=NAVY, bold=True)
    for label, value in control_rows:
        cells = table.add_row().cells
        add_inline(cells[0].paragraphs[0], label, size=9, color=NAVY, bold=True)
        add_inline(cells[1].paragraphs[0], value, size=9, color=INK)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [1900, 7460])

    doc.add_heading("目录", level=1)
    p = doc.add_paragraph()
    add_field(p, ' TOC \\o "1-3" \\h \\z \\u ', "目录将在 Word/LibreOffice 中更新")


def visual_length(text: str) -> int:
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


def choose_widths(headers: list[str], rows: list[list[str]]) -> list[int]:
    count = len(headers)
    if count == 1:
        return [CONTENT_WIDTH_DXA]
    if count == 2:
        first = max(1650, min(3100, 700 + max(visual_length(row[0]) for row in [headers] + rows) * 55))
        return [first, CONTENT_WIDTH_DXA - first]
    if count == 3:
        combined = [headers] + rows
        maxlens = [max(visual_length(row[i]) for row in combined) for i in range(3)]
        minimums = [1300, 1900, 1900]
        weights = [max(1.0, min(4.0, length ** 0.5)) for length in maxlens]
        remaining = CONTENT_WIDTH_DXA - sum(minimums)
        widths = [minimums[i] + int(remaining * weights[i] / sum(weights)) for i in range(3)]
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
        return widths
    minimum = 1100
    widths = [minimum] * count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def use_record_cards(headers: list[str], rows: list[list[str]]) -> bool:
    if len(headers) >= 5:
        return True
    if len(headers) == 4:
        body_lengths = [visual_length(cell) for row in rows for cell in row]
        return max(body_lengths, default=0) > 90 or sum(body_lengths) / max(1, len(body_lengths)) > 42
    return False


def add_record_cards(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    for index, row in enumerate(rows, start=1):
        card = doc.add_paragraph()
        card.paragraph_format.space_before = Pt(5)
        card.paragraph_format.space_after = Pt(3)
        card.paragraph_format.keep_with_next = True
        shade(card, LIGHT_BLUE if index % 2 else LIGHT_GRAY)
        set_paragraph_border(card, side="left", color=BLUE, size=12, space=8)
        title = row[0] if row else f"记录 {index}"
        add_inline(card, title, size=10, color=NAVY, bold=True)
        for label, value in zip(headers[1:], row[1:]):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.12
            label_run = p.add_run(f"{label}：")
            set_run_font(label_run, size=8.6, color=DARK_BLUE, bold=True)
            add_inline(p, value, size=8.6, color=INK)


def add_markdown_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    headers = [normalize_inline(x) for x in headers]
    rows = [[normalize_inline(x) for x in row] for row in rows]
    if use_record_cards(headers, rows):
        add_record_cards(doc, headers, rows)
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        add_inline(cell.paragraphs[0], header, size=8.5, color=NAVY, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_table_row_split(table.rows[-1])
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            if row_index % 2:
                shade(cell, "FAFBFC")
            add_inline(cell.paragraphs[0], value, size=8.2, color=INK,
                       bold=(col_index == 0 and len(headers) <= 3))
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.08
            if col_index == 0 and visual_length(value) < 20:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, choose_widths(headers, rows))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [part.strip().replace("\\|", "|") for part in line.split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_code_block(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    # Keep command/flow diagrams out of the running header when Word
    # repaginates a block across pages during DOCX-to-PDF verification.
    p.paragraph_format.keep_together = True
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.05
    shade(p, "EEF1F4")
    set_paragraph_border(p, side="left", color=MUTED, size=8, space=8)
    for idx, line in enumerate(code.rstrip().splitlines()):
        run = p.add_run(line)
        set_run_font(run, size=8.3, color="263238", mono=True)
        if idx != len(code.rstrip().splitlines()) - 1:
            run.add_break()


def add_callout(doc: Document, lines: Iterable[str]) -> None:
    text = " ".join(line.lstrip("> ").strip() for line in lines)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    shade(p, CALLOUT)
    set_paragraph_border(p, side="left", color=BLUE, size=14, space=10)
    add_inline(p, text, size=9.5, color=INK)


def add_figure(doc: Document) -> None:
    if not SCREENSHOT.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    inline_shape = run.add_picture(str(SCREENSHOT), width=Inches(6.45))
    docpr = inline_shape._inline.docPr
    docpr.set("name", "LI3D texture editor module overview")
    docpr.set("descr", "LI3D 贴图编辑器全界面，用于对应 UI-01 至 UI-12 模块索引")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run("图 1  LI3D 贴图编辑器界面与 UI 模块审计基准")
    set_run_font(run, size=8.5, color=MUTED, italic=True)


def should_page_break_heading(text: str, level: int) -> bool:
    if level != 1:
        return False
    match = re.match(r"(\d+)\.", text.strip())
    if not match:
        return False
    # The dynamic Word TOC already ends on a page boundary. Forcing chapter 1 to
    # break again creates a fully blank page after Word expands the TOC. Avoid
    # forced breaks for chapters 6 and 16 so short continuation text from the
    # previous chapter can share the page.
    return int(match.group(1)) in {3, 5, 8, 14, 17, 18, 19, 20, 21, 22}


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    if should_page_break_heading(text, level):
        p.paragraph_format.page_break_before = True
    if text.startswith("14. 编辑器界面模块总图"):
        add_figure(doc)


def render_markdown(doc: Document, markdown: str, bullet_num_id: int, decimal_num_id: int) -> None:
    lines = markdown.splitlines()
    index = 0
    paragraph_buffer: list[str] = []
    seen_main_heading = False

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        if text:
            p = doc.add_paragraph()
            add_inline(p, text, size=11, color=INK)
        paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not seen_main_heading:
            if stripped.startswith("## 1."):
                seen_main_heading = True
            else:
                index += 1
                continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1)) - 1
            add_heading(doc, heading_match.group(2).strip(), level)
            index += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(doc, "\n".join(code_lines))
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index])
                index += 1
            add_callout(doc, quote_lines)
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_separator_row(lines[index + 1]):
            flush_paragraph()
            headers = split_table_row(line)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = split_table_row(lines[index])
                if len(row) < len(headers):
                    row += [""] * (len(headers) - len(row))
                rows.append(row[: len(headers)])
                index += 1
            add_markdown_table(doc, headers, rows)
            continue

        bullet_match = re.match(r"^(\s*)-\s+(.+)$", line)
        if bullet_match:
            flush_paragraph()
            level = min(2, len(bullet_match.group(1).replace("\t", "  ")) // 2)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_numbering(p, bullet_num_id, level)
            add_inline(p, bullet_match.group(2), size=10.5, color=INK)
            index += 1
            continue

        number_match = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if number_match:
            flush_paragraph()
            level = min(2, len(number_match.group(1).replace("\t", "  ")) // 2)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_numbering(p, decimal_num_id, level)
            add_inline(p, number_match.group(2), size=10.5, color=INK)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1
    flush_paragraph()


def add_document_metadata(doc: Document) -> None:
    props = doc.core_properties
    props.title = "LI3D 系统模块与算法维护手册"
    props.subject = "模块边界、投影、UV 合成、局部重绘、图层、保存、版本和变更控制"
    props.author = "LI3D Engineering"
    props.keywords = "LI3D, 3D Texture, Projection, UV, Local Repaint, Maintenance"
    props.comments = f"Normative maintenance baseline v{DOCUMENT_VERSION}; source {SOURCE.name}"


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
        configure_header_footer(section)
    set_update_fields_on_open(doc)
    bullet_num_id = create_numbering(doc, ordered=False)
    decimal_num_id = create_numbering(doc, ordered=True)
    add_document_metadata(doc)
    add_cover(doc)
    add_front_matter(doc)
    render_markdown(doc, markdown, bullet_num_id, decimal_num_id)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Preset={PRESET_NAME}; Header={HEADER_PATTERN}; Source={SOURCE}")


if __name__ == "__main__":
    main()
